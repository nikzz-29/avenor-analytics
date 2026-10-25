import csv
import hashlib
import io
import json
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import PurePath
from typing import Literal
from xml.etree.ElementTree import ParseError

import openpyxl
import polars as pl
from defusedxml.common import DefusedXmlException
from fastapi import APIRouter, File, Form, HTTPException, Query, Response, UploadFile
from pydantic import BaseModel, Field, ValidationError, model_validator
from sqlalchemy import delete, insert, select
from starlette.concurrency import run_in_threadpool

from app.core.security import DB, Membership
from app.db.models import AuditLog, Dataset, ImportJob, Sale, Workspace
from app.modules.analytics import csv_response
from app.modules.schemas import (
    DatasetOut,
    ImportErrorOut,
    ImportJobOut,
    ImportOut,
    PreviewOut,
    SaleOut,
)

router = APIRouter(prefix="/workspaces/{workspace_id}", tags=["Data sources"])
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 100_000
REQUIRED = {"order_id", "order_date", "customer_id", "product_name", "quantity", "unit_price"}


class SaleInput(BaseModel):
    order_id: str = Field(min_length=1, max_length=200)
    order_date: date
    customer_id: str = Field(min_length=1, max_length=200)
    product_name: str = Field(min_length=1, max_length=200)
    category: str = Field(default="Uncategorized", min_length=1, max_length=200)
    channel: str = Field(default="Unknown", min_length=1, max_length=200)
    status: Literal["paid", "pending", "cancelled", "refunded"] = "paid"
    quantity: int = Field(ge=1, le=100000)
    unit_price: Decimal = Field(ge=0, max_digits=12, decimal_places=2)
    discount: Decimal = Field(default=Decimal(0), ge=0, max_digits=14, decimal_places=2)
    refund: Decimal = Field(default=Decimal(0), ge=0, max_digits=14, decimal_places=2)

    @model_validator(mode="after")
    def money(self) -> "SaleInput":
        gross = self.quantity * self.unit_price
        if gross >= Decimal("1000000000000"):
            raise ValueError("Line amount exceeds supported range")
        if self.discount + self.refund > gross:
            raise ValueError("Discount plus refund exceeds gross line amount")
        if self.refund and self.status != "refunded":
            raise ValueError("Refund amount requires refunded status")
        return self


def parse_file(content: bytes, filename: str) -> tuple[list[str], list[dict[str, str]]]:
    if not content or len(content) > MAX_BYTES:
        raise ValueError("Choose a non-empty file up to 10 MiB")
    suffix = PurePath(filename).suffix.lower()
    raw_rows: list[list[object]] = []
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            if "\x00" in text:
                raise ValueError("CSV must not contain binary NUL bytes")
            for physical_row, row in enumerate(csv.reader(io.StringIO(text), dialect)):
                if physical_row > MAX_ROWS:
                    raise ValueError("Maximum 100,000 rows per import")
                if any(cell.strip() for cell in row):
                    raw_rows.append(list(row))
                if len(raw_rows) > MAX_ROWS + 1:
                    raise ValueError("Maximum 100,000 rows per import")
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ValueError("Expected UTF-8 CSV with a header row") from exc
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Uncompressed spreadsheet is too large")
                if len(archive.infolist()) > 1000:
                    raise ValueError("Spreadsheet archive contains too many entries")
            workbook = openpyxl.load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            sheet = workbook.active
            if sheet is None:
                raise ValueError("Spreadsheet has no worksheet")
            if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                raise ValueError("Maximum 100,000 worksheet rows per import")
            if sheet.max_column and sheet.max_column > 50:
                raise ValueError("Maximum 50 columns")
            for cells in sheet.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Spreadsheet formulas are not supported; export values first")
                values: list[object] = [cell.value for cell in cells]
                if any(v is not None for v in values):
                    raw_rows.append(values)
                if len(raw_rows) > MAX_ROWS + 1:
                    raise ValueError("Maximum 100,000 rows per import")
            workbook.close()
        except (zipfile.BadZipFile, KeyError, OSError, ParseError, DefusedXmlException) as exc:
            raise ValueError("Invalid XLSX file") from exc
    else:
        raise ValueError("Only CSV and XLSX files are supported")
    if len(raw_rows) < 2:
        raise ValueError("Include a header and at least one data row")
    columns = [str(value or "").strip() for value in raw_rows[0]]
    if (
        len(columns) > 50
        or any(not c or len(c) > 200 for c in columns)
        or len(set(columns)) != len(columns)
    ):
        raise ValueError("Use up to 50 unique, non-empty column names (maximum 200 characters)")
    result: list[dict[str, str]] = []
    for i, values in enumerate(raw_rows[1:], start=2):
        if len(values) != len(columns):
            raise ValueError(f"Row {i} has {len(values)} columns; expected {len(columns)}")
        normalized = [
            v.date().isoformat()
            if isinstance(v, datetime)
            else str(v).strip()
            if v is not None
            else ""
            for v in values
        ]
        if any(len(v) > 2000 for v in normalized):
            raise ValueError(f"Row {i} contains a cell longer than 2,000 characters")
        result.append(dict(zip(columns, normalized, strict=True)))
    return columns, result


async def read_file(file: UploadFile) -> tuple[bytes, list[str], list[dict[str, str]]]:
    content = await file.read(MAX_BYTES + 1)
    try:
        columns, rows = await run_in_threadpool(parse_file, content, file.filename or "")
    except (ValueError, TypeError) as exc:
        raise HTTPException(422, str(exc)) from exc
    return content, columns, rows


@router.post("/datasets/preview", response_model=PreviewOut)
async def preview(member: Membership, file: UploadFile = File(...)) -> PreviewOut:
    _, columns, rows = await read_file(file)
    inferred: dict[str, str] = {}
    for column in columns:
        sample = [r[column] for r in rows[:100] if r[column]]
        series = pl.Series(sample, dtype=pl.String)
        if not sample:
            kind = "empty"
        elif series.cast(pl.Float64, strict=False).null_count() == 0:
            kind = "number"
        else:
            try:
                kind = "date" if series.str.to_date(strict=False).null_count() == 0 else "text"
            except pl.exceptions.PolarsError:
                kind = "text"
        inferred[column] = kind
    return PreviewOut(columns=columns, rows=rows[:8], total_rows=len(rows), inferred_types=inferred)


def validate_rows(
    rows: list[dict[str, str]], mapping: dict[str, str]
) -> tuple[list[dict[str, object]], list[ImportErrorOut], int]:
    valid: list[dict[str, object]] = []
    errors: list[ImportErrorOut] = []
    seen: set[tuple[str, str]] = set()
    orders: dict[str, tuple[date, str, str, str]] = {}
    invalid = 0
    for index, row in enumerate(rows, start=2):
        mapped = {
            field: row[source]
            for field, source in mapping.items()
            if source in row and row[source] != ""
        }
        try:
            item = SaleInput.model_validate(mapped)
            key = (item.order_id, item.product_name)
            if key in seen:
                raise ValueError("Duplicate order/product line in file")
            identity = (item.order_date, item.customer_id, item.status, item.channel)
            if item.order_id in orders and orders[item.order_id] != identity:
                raise ValueError("Lines in one order must share date, customer, status and channel")
            orders[item.order_id] = identity
            seen.add(key)
            valid.append(item.model_dump())
        except (ValidationError, ValueError) as exc:
            invalid += 1
            if len(errors) < 1000:
                if isinstance(exc, ValidationError):
                    for err in exc.errors()[:3]:
                        if len(errors) < 1000:
                            errors.append(
                                ImportErrorOut(
                                    row=index,
                                    column=".".join(str(v) for v in err["loc"]) or "row",
                                    message=err["msg"],
                                )
                            )
                else:
                    errors.append(ImportErrorOut(row=index, column="order_id", message=str(exc)))
    return valid, errors, invalid


@router.post("/datasets/import", response_model=ImportOut)
async def import_dataset(
    member: Membership,
    db: DB,
    file: UploadFile = File(...),
    mapping: str = Form("{}"),
    name: str = Form("Imported sales"),
    mode: Literal["create", "replace"] = Form("create"),
    dataset_id: str | None = Form(None),
    validate_only: bool = Form(False),
) -> ImportOut:
    if not name.strip() or len(name) > 200:
        raise HTTPException(422, "Dataset name must be 1–200 characters")
    content, columns, rows = await read_file(file)
    try:
        parsed = json.loads(mapping)
        if not isinstance(parsed, dict) or any(
            not isinstance(k, str) or not isinstance(v, str) for k, v in parsed.items()
        ):
            raise ValueError("Mapping must be a field-to-column object")
        parsed = parsed or {c: c for c in columns if c in SaleInput.model_fields}
        if not REQUIRED.issubset(parsed) or any(
            k not in SaleInput.model_fields or v not in columns for k, v in parsed.items()
        ):
            raise ValueError("Map each required field to an existing column")
    except (ValueError, TypeError) as exc:
        raise HTTPException(422, str(exc)) from exc
    valid, errors, rejected = await run_in_threadpool(validate_rows, rows, parsed)
    if validate_only:
        return ImportOut(
            id="validation",
            status="failed" if rejected else "completed",
            imported=0,
            rejected=rejected,
            errors=errors,
        )
    await db.scalar(select(Workspace).where(Workspace.id == member.workspace_id).with_for_update())
    target: Dataset | None = None
    if mode == "replace":
        target = await db.scalar(
            select(Dataset).where(
                Dataset.id == dataset_id, Dataset.workspace_id == member.workspace_id
            )
        )
        if not target:
            raise HTTPException(404, "Choose a dataset to replace")
    fingerprint = hashlib.sha256(content).hexdigest()
    existing = await db.scalar(
        select(Dataset).where(
            Dataset.workspace_id == member.workspace_id, Dataset.fingerprint == fingerprint
        )
    )
    if existing and (not target or target.id != existing.id):
        raise HTTPException(409, "This file has already been imported")
    if not rejected:
        order_ids = list({str(row["order_id"]) for row in valid})
        for offset in range(0, len(order_ids), 1000):
            stmt = select(Sale.order_id).where(
                Sale.workspace_id == member.workspace_id,
                Sale.order_id.in_(order_ids[offset : offset + 1000]),
            )
            if target:
                stmt = stmt.where(Sale.dataset_id != target.id)
            duplicate = await db.scalar(stmt.limit(1))
            if duplicate:
                errors.append(
                    ImportErrorOut(
                        row=0,
                        column="order_id",
                        message=f"Order {duplicate} already belongs to another dataset. Use replace to avoid double counting.",
                    )
                )
                rejected = len(rows)
                break
    job = ImportJob(
        workspace_id=member.workspace_id,
        filename=PurePath(file.filename or "upload").name[:200],
        status="failed" if rejected else "completed",
        imported=0 if rejected else len(valid),
        rejected=rejected,
        errors=[e.model_dump() for e in errors],
    )
    db.add(job)
    await db.flush()
    if not rejected:
        if target:
            await db.execute(delete(Sale).where(Sale.dataset_id == target.id))
            target.fingerprint, target.row_count, target.name = (
                fingerprint,
                len(valid),
                name.strip(),
            )
            dataset = target
        else:
            dataset = Dataset(
                workspace_id=member.workspace_id,
                name=name.strip(),
                fingerprint=fingerprint,
                row_count=len(valid),
                columns=[
                    {
                        "name": field,
                        "type": "number"
                        if field in ("quantity", "unit_price", "discount", "refund")
                        else "date"
                        if field == "order_date"
                        else "text",
                    }
                    for field in parsed
                ],
            )
            db.add(dataset)
            await db.flush()
        for offset in range(0, len(valid), 1000):
            batch = [
                {**row, "workspace_id": member.workspace_id, "dataset_id": dataset.id}
                for row in valid[offset : offset + 1000]
            ]
            await db.execute(insert(Sale), batch)
        db.add(
            AuditLog(
                workspace_id=member.workspace_id,
                action=f"Imported {len(valid):,} lines · {dataset.name}",
            )
        )
    return ImportOut(
        id=job.id, status=job.status, imported=job.imported, rejected=job.rejected, errors=errors
    )


@router.get("/datasets", response_model=list[DatasetOut])
async def datasets(member: Membership, db: DB) -> list[Dataset]:
    return list(
        await db.scalars(
            select(Dataset)
            .where(Dataset.workspace_id == member.workspace_id)
            .order_by(Dataset.created_at.desc())
        )
    )


@router.get("/datasets/{dataset_id}/preview", response_model=list[SaleOut])
async def dataset_preview(dataset_id: str, member: Membership, db: DB) -> list[Sale]:
    return list(
        await db.scalars(
            select(Sale)
            .where(Sale.workspace_id == member.workspace_id, Sale.dataset_id == dataset_id)
            .limit(20)
        )
    )


@router.delete("/datasets/{dataset_id}", response_model=dict[str, str])
async def remove_dataset(dataset_id: str, member: Membership, db: DB) -> dict[str, str]:
    dataset = await db.scalar(
        select(Dataset).where(Dataset.id == dataset_id, Dataset.workspace_id == member.workspace_id)
    )
    if not dataset:
        raise HTTPException(404, "Dataset not found")
    await db.delete(dataset)
    db.add(AuditLog(workspace_id=member.workspace_id, action=f"Deleted dataset · {dataset.name}"))
    return {"message": "Dataset deleted"}


@router.get("/imports", response_model=list[ImportJobOut])
async def jobs(
    member: Membership, db: DB, limit: int = Query(default=20, ge=1, le=100)
) -> list[ImportJob]:
    return list(
        await db.scalars(
            select(ImportJob)
            .where(ImportJob.workspace_id == member.workspace_id)
            .order_by(ImportJob.created_at.desc())
            .limit(limit)
        )
    )


@router.get("/imports/{job_id}/errors")
async def errors_csv(job_id: str, member: Membership, db: DB) -> Response:
    job = await db.scalar(
        select(ImportJob).where(
            ImportJob.id == job_id, ImportJob.workspace_id == member.workspace_id
        )
    )
    if not job:
        raise HTTPException(404, "Import not found")
    return csv_response(
        [
            ["row", "column", "message"],
            *[[str(e["row"]), str(e["column"]), str(e["message"])] for e in job.errors],
        ],
        "import-errors.csv",
    )
