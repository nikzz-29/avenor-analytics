"""Regression coverage for edge cases found during the project audit."""

from decimal import Decimal

import httpx
import pytest

from app.modules.imports import parse_file, validate_rows
from tests.conftest import workspace
from tests.test_api import CSV, upload


@pytest.mark.parametrize("endpoint", ["sales", "entities/products", "entities/customers"])
async def test_reversed_get_range_is_validation_error(
    client: httpx.AsyncClient, endpoint: str
) -> None:
    base = await workspace(client)
    response = await client.get(base + f"/analytics/{endpoint}?start=2026-02-01&end=2026-01-01")
    assert response.status_code == 422


async def test_overview_does_not_overflow_previous_period(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    response = await client.get(base + "/analytics/overview?start=0001-01-01&end=0001-01-02")
    assert response.status_code == 422


async def test_large_numeric_filter_is_rejected_before_database(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    response = await client.post(
        base + "/analytics/query",
        json={"filters": [{"field": "unit_price", "op": "gt", "values": ["1e1000000"]}]},
    )
    assert response.status_code == 422


async def test_nul_name_is_a_validation_error(client: httpx.AsyncClient) -> None:
    await workspace(client)
    response = await client.post("/api/v1/workspaces", json={"name": "Bad\x00name"})
    assert response.status_code == 422


async def test_patch_workspace_preserves_omitted_default(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    dashboard = (await client.post(base + "/dashboards", json={"name": "Default"})).json()
    assert (
        await client.patch(base, json={"name": "Original", "default_dashboard_id": dashboard["id"]})
    ).status_code == 200
    result = await client.patch(base, json={"name": "Renamed"})
    assert result.status_code == 200
    assert result.json()["default_dashboard_id"] == dashboard["id"]


async def test_replace_refreshes_dataset_schema(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    await upload(client, base)
    dataset = (await client.get(base + "/datasets")).json()[0]
    replacement = b"order_id,order_date,customer_id,product_name,quantity,unit_price\nNEW,2026-01-01,A,Chair,1,15.50\n"
    response = await upload(client, base, replacement, mode="replace", dataset_id=dataset["id"])
    assert response.json()["status"] == "completed"
    updated = (await client.get(base + "/datasets")).json()[0]
    assert {c["name"] for c in updated["columns"]} == {
        "order_id",
        "order_date",
        "customer_id",
        "product_name",
        "quantity",
        "unit_price",
    }


async def test_missing_dataset_preview_is_not_silent_empty(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    response = await client.get(base + "/datasets/00000000-0000-0000-0000-000000000000/preview")
    assert response.status_code == 404


async def test_report_query_aov_undefined_is_null_not_zero(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    await upload(client, base)
    response = await client.post(
        base + "/analytics/query", json={"metric": "aov", "dimension": "status"}
    )
    assert response.status_code == 200
    cancelled = next(p for p in response.json()["points"] if p["label"] == "cancelled")
    assert cancelled["value"] is None


def test_import_date_requires_iso_date_not_unix_timestamp() -> None:
    columns, rows = parse_file(CSV, "sales.csv")
    rows[0]["order_date"] = "0"
    _, errors, rejected = validate_rows(rows, {c: c for c in columns})
    assert rejected == 1 and errors[0].column == "order_date"


def test_blank_lines_preserve_validation_row_numbers() -> None:
    columns, rows = parse_file(
        b"order_id,order_date,customer_id,product_name,quantity,unit_price\n\nA,2025-01-01,C,Lamp,-1,10\n",
        "sales.csv",
    )
    _, errors, rejected = validate_rows(rows, {c: c for c in columns})
    assert rejected == 1 and errors[0].row == 3


async def test_comparison_projects_previous_days_into_current_buckets(
    client: httpx.AsyncClient,
) -> None:
    base = await workspace(client)
    data = b"order_id,order_date,customer_id,product_name,quantity,unit_price\nP1,2025-01-29,C,Lamp,1,20\nP2,2025-02-28,C,Lamp,1,80\nC1,2025-03-31,C,Lamp,1,50\n"
    await upload(client, base, data)
    response = await client.get(base + "/analytics/overview?start=2025-03-01&end=2025-03-31")
    assert response.status_code == 200
    value = response.json()
    assert value["trend"] == [{"label": "2025-03-01", "value": "50.00"}]
    assert value["comparison"] == [{"label": "2025-03-01", "value": "100.00"}]


async def test_daily_overview_does_not_silently_stop_at_100_points(
    client: httpx.AsyncClient,
) -> None:
    base = await workspace(client)
    await upload(client, base)
    response = await client.get(
        base + "/analytics/overview?start=2025-01-01&end=2025-12-31&grouping=day"
    )
    assert response.status_code == 200
    data = response.json()
    assert len(data["trend"]) == 365 and len(data["comparison"]) == 365
    assert data["trend"][-1]["label"] == "2025-12-31"
    assert sum(Decimal(p["value"]) for p in data["trend"]) == 250


async def test_oversized_chart_period_is_explicit(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    response = await client.get(
        base + "/analytics/overview?start=2020-01-01&end=2025-12-31&grouping=day"
    )
    assert response.status_code == 422


async def test_workspace_default_period_follows_imported_data(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    response = await upload(client, base, CSV.replace(b"2025-", b"2026-"))
    assert response.json()["status"] == "completed"
    period = await client.get(base + "/analytics/period")
    assert period.status_code == 200
    assert period.json() == {"start": "2026-01-01", "end": "2026-12-31"}


async def test_concurrent_demo_population_is_serialized(client: httpx.AsyncClient) -> None:
    import asyncio

    base = await workspace(client)
    results = await asyncio.gather(client.post(base + "/demo"), client.post(base + "/demo"))
    assert sorted(r.status_code for r in results) == [200, 409]
    datasets = (await client.get(base + "/datasets")).json()
    assert len(datasets) == 1 and datasets[0]["row_count"] == 2548


async def test_concurrent_imports_do_not_duplicate_facts(client: httpx.AsyncClient) -> None:
    import asyncio

    base = await workspace(client)
    results = await asyncio.gather(upload(client, base), upload(client, base))
    assert sorted(r.status_code for r in results) == [200, 409]
    overview = (await client.get(base + "/analytics/overview")).json()
    assert Decimal(overview["current"]["revenue"]) == 250


async def test_query_bounds_and_control_characters_are_rejected(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    for path in (
        "/analytics/sales?status=unknown",
        "/analytics/entities/customers?page=" + ("9" * 50),
        "/analytics/sales?search=%00",
        "/datasets/%00/preview",
    ):
        response = await client.get(base + path)
        assert response.status_code == 422, path
    response = await client.post(
        base + "/analytics/query",
        json={"filters": [{"field": "category", "op": "eq", "values": ["bad\x00value"]}]},
    )
    assert response.status_code == 422


async def test_failed_replace_preserves_original_data(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    await upload(client, base)
    dataset = (await client.get(base + "/datasets")).json()[0]
    invalid = CSV.replace(b"2,50", b"-2,50")
    response = await upload(client, base, invalid, mode="replace", dataset_id=dataset["id"])
    assert response.json()["status"] == "failed"
    assert (await client.get(base + "/datasets")).json()[0] == dataset
    assert (
        Decimal((await client.get(base + "/analytics/overview")).json()["current"]["revenue"])
        == 250
    )


def test_unclosed_csv_quote_is_rejected() -> None:
    with pytest.raises(ValueError):
        parse_file(b'order_id,product_name\nA,"Unclosed\n', "sales.csv")


def test_xlsx_without_dimensions_still_checks_actual_column_count() -> None:
    import io
    import re
    import zipfile

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append([f"Column{i}" for i in range(51)])
    sheet.append(list(range(51)))
    buffer = io.BytesIO()
    workbook.save(buffer)
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(buffer.getvalue())) as source,
        zipfile.ZipFile(output, "w") as target,
    ):
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = re.sub(rb"<dimension[^>]*/>", b"", content)
            target.writestr(item, content)
    with pytest.raises(ValueError, match="50 columns"):
        parse_file(output.getvalue(), "sales.xlsx")


async def test_foreign_children_are_scoped_even_for_a_member_of_both_workspaces(
    client: httpx.AsyncClient,
) -> None:
    base = await workspace(client)
    second = (await client.post("/api/v1/workspaces", json={"name": "Second"})).json()
    other = "/api/v1/workspaces/" + second["id"]
    await upload(client, other)
    dataset = (await client.get(other + "/datasets")).json()[0]
    board = (await client.post(other + "/dashboards", json={"name": "Private view"})).json()
    report = (
        await client.post(other + "/reports", json={"name": "Private report", "query": {}})
    ).json()
    assert (await client.get(base + "/datasets/" + dataset["id"] + "/preview")).status_code == 404
    assert (await client.delete(base + "/datasets/" + dataset["id"])).status_code == 404
    assert (
        await client.post(base + "/analytics/query", json={"dataset_id": dataset["id"]})
    ).status_code == 404
    assert (await client.get(base + "/dashboards/" + board["id"])).status_code == 404
    assert (
        await client.patch(
            base + "/dashboards/" + board["id"], json={"name": "Wrong", "revision": 1}
        )
    ).status_code == 404
    assert (await client.delete(base + "/dashboards/" + board["id"])).status_code == 404
    assert (
        await client.patch(base, json={"name": "First", "default_dashboard_id": board["id"]})
    ).status_code == 404
    assert (await client.delete(base + "/reports/" + report["id"])).status_code == 404
    assert len((await client.get(other + "/datasets")).json()) == 1
    assert (await client.get(other + "/dashboards/" + board["id"])).json()["name"] == "Private view"


@pytest.mark.parametrize("role", ["editor", "admin"])
async def test_editor_and_admin_cannot_escalate_to_owner(
    client: httpx.AsyncClient, role: str
) -> None:
    from tests.conftest import signup

    base = await workspace(client)
    owner = dict(client.cookies)
    client.cookies.clear()
    await signup(client, "member@example.com")
    member = dict(client.cookies)
    client.cookies.clear()
    client.cookies.update(owner)
    client.headers["x-csrf-token"] = client.cookies["avenor_csrf"]
    assert (
        await client.put(base + "/members", json={"email": "member@example.com", "role": role})
    ).status_code == 200
    owner_id = next(
        m["user_id"] for m in (await client.get(base + "/members")).json() if m["role"] == "owner"
    )
    client.cookies.clear()
    client.cookies.update(member)
    client.headers["x-csrf-token"] = client.cookies["avenor_csrf"]
    assert (await client.post(base + "/dashboards", json={"name": "Editable"})).status_code == 201
    assert (await client.delete(base + "/members/" + owner_id)).status_code == 403
    assert (
        await client.put(base + "/members", json={"email": "member@example.com", "role": "owner"})
    ).status_code == 422
    assert (
        await client.put(base + "/members", json={"email": "member@example.com", "role": "admin"})
    ).status_code == 403


async def test_empty_additive_metrics_and_undefined_averages_are_distinct(
    client: httpx.AsyncClient,
) -> None:
    base = await workspace(client)
    for metric in ("revenue", "orders", "units", "customers", "sum", "count", "distinct_count"):
        response = await client.post(base + "/analytics/query", json={"metric": metric})
        assert response.status_code == 200
        assert Decimal(response.json()["total"]) == 0
        assert response.json()["points"] == []
    for metric in ("aov", "average", "minimum", "maximum"):
        response = await client.post(base + "/analytics/query", json={"metric": metric})
        assert response.status_code == 200
        assert response.json()["total"] is None


async def test_overview_remains_consistent_when_an_import_commits_mid_request(
    client: httpx.AsyncClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    import asyncio

    from app.modules import analytics

    base = await workspace(client)
    await upload(client, base)
    current_read = asyncio.Event()
    continue_reader = asyncio.Event()
    original = analytics.calculate_kpi
    first = True

    async def pause_after_current(*args, **kwargs):
        nonlocal first
        result = await original(*args, **kwargs)
        if first:
            first = False
            current_read.set()
            await continue_reader.wait()
        return result

    monkeypatch.setattr(analytics, "calculate_kpi", pause_after_current)
    reader = asyncio.create_task(client.get(base + "/analytics/overview"))
    await asyncio.wait_for(current_read.wait(), timeout=5)
    additional = b"order_id,order_date,customer_id,product_name,quantity,unit_price\nNEW,2025-06-04,New customer,Chair,1,75\n"
    writer = asyncio.create_task(upload(client, base, additional))
    try:
        # Either the writer commits, or it correctly waits for the read transaction.
        # No performance assertion depends on this scheduling window.
        await asyncio.wait({writer}, timeout=1)
    finally:
        continue_reader.set()
    response, imported = await asyncio.gather(reader, writer)
    assert imported.status_code == 200 and imported.json()["status"] == "completed"
    assert response.status_code == 200
    data = response.json()
    assert Decimal(data["current"]["revenue"]) == sum(Decimal(p["value"]) for p in data["trend"])


async def test_workspace_patch_changes_only_explicit_fields(client: httpx.AsyncClient) -> None:
    base = await workspace(client)
    board = (await client.post(base + "/dashboards", json={"name": "Default"})).json()
    assert (await client.patch(base, json={"name": "Latest name"})).status_code == 200
    selected = await client.patch(base, json={"default_dashboard_id": board["id"]})
    assert selected.status_code == 200
    assert selected.json()["name"] == "Latest name"
    assert selected.json()["default_dashboard_id"] == board["id"]
    assert (await client.patch(base, json={"name": None})).status_code == 422


def test_active_excel_chartsheet_is_a_clear_validation_error() -> None:
    import io

    import openpyxl
    from openpyxl.chart import BarChart, Reference

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Order", "Date", "Customer", "Product", "Qty", "Price"])
    sheet.append(["A", "2025-01-01", "C", "Lamp", 1, 10])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=6, min_row=1, max_row=2), titles_from_data=True)
    workbook.create_chartsheet("Chart").add_chart(chart)
    workbook.active = 1
    output = io.BytesIO()
    workbook.save(output)
    with pytest.raises(ValueError, match="worksheet"):
        parse_file(output.getvalue(), "sales.xlsx")


def test_password_protected_archive_is_rejected_cleanly() -> None:
    import io
    import zipfile

    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
    data = bytearray(output.getvalue())
    for signature, offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        position = data.index(signature) + offset
        data[position : position + 2] = (
            int.from_bytes(data[position : position + 2], "little") | 1
        ).to_bytes(2, "little")
    with pytest.raises(ValueError, match="XLSX"):
        parse_file(bytes(data), "encrypted.xlsx")
